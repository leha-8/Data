# -*- coding: utf-8 -*-
import io
import sys
import socket
from datetime import datetime
from flask import Flask, render_template, request, jsonify, send_file
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import models
import scanner
import os

if getattr(sys, 'frozen', False):
    template_folder = os.path.join(sys._MEIPASS, 'templates') if os.path.exists(os.path.join(sys._MEIPASS, 'templates')) else 'templates'
    static_folder = os.path.join(sys._MEIPASS, 'static') if os.path.exists(os.path.join(sys._MEIPASS, 'static')) else 'static'
    app = Flask(__name__, template_folder=template_folder, static_folder=static_folder)
else:
    app = Flask(__name__)

# ... (các đoạn code phía dưới giữ nguyên)

# Khởi tạo CSDL và migration khi khởi động ứng dụng
models.init_db()

LOCAL_IP = scanner.get_local_ip()

@app.route("/")
def index():
    return render_template("index.html", local_ip=LOCAL_IP)

# --- API DASHBOARD ---
@app.route("/api/dashboard", methods=["GET"])
def api_dashboard():
    stats = models.get_dashboard_stats()
    return jsonify(stats)

# --- API SITES (CƠ SỞ / CHI NHÁNH) ---
@app.route("/api/sites", methods=["GET"])
def api_get_sites():
    sites = [
        "HBC", "HBC-BH", "Q12", "Q12-ĐNTH02", "Q12-HHG", 
        "Q12-LTR", "Q12-TTN17", "Q12-TTN6", "Q2", "Q2-111LDC", 
        "Q2-BAO2-LDC", "Q2-PC", "Q7-ECOGREEN", "Q9-FNL", "QTB-NSS", 
        "TDTT", "TPTD-DK", "TPTD-NTD", "TPTD-TML"
    ]
    return jsonify(sites)

# --- API NETWORK SCANNER (QUÉT MẠNG LAN) ---
@app.route("/api/network-info", methods=["GET"])
def api_network_info():
    """Lấy thông tin mạng hiện tại: IP card mạng, Subnet gợi ý và các cơ sở cấu hình sẵn."""
    info = scanner.get_network_info()
    return jsonify(info)

@app.route("/api/scan-lan", methods=["POST"])
def api_scan_lan():
    """
    Quét dải mạng LAN linh hoạt theo yêu cầu:
    Nhận subnet_ip từ client (tự động phát hiện, cơ sở định sẵn hoặc tùy chỉnh).
    """
    data = request.json or {}
    
    # Lấy tên cơ sở mặc định nếu client không truyền lên
    all_sites = models.get_all_sites()
    default_site = all_sites[0] if all_sites else "HBC"
    
    subnet = data.get("subnet_ip") or data.get("subnet") or ""
    site_name = data.get("site_name", default_site).strip() or default_site
    
    if not subnet:
        net_info = scanner.get_network_info()
        subnet = net_info["suggested_subnet"]

    try:
        inventory_items = models.get_all_items()
        # Truyền thêm site_name để scanner biết nếu là TPTD-TML thì quét thêm dải 200.x
        result = scanner.scan_subnet(subnet, inventory_items=inventory_items, site_name=site_name)
        result["target_site"] = site_name
        return jsonify(result)
    except ValueError as ve:
        return jsonify({"error": str(ve)}), 400
    except Exception as e:
        return jsonify({"error": f"Lỗi trong quá trình quét dải mạng: {str(e)}"}), 500

@app.route("/api/items/quick-add-from-scan", methods=["POST"])
def api_quick_add_from_scan():
    """Thêm thiết bị vừa phát hiện qua Quét LAN vào danh mục quản lý kho."""
    data = request.json or {}
    name = data.get("name", "").strip()
    if not name:
        return jsonify({"error": "Tên thiết bị không được để trống"}), 400

    try:
        item_payload = {
            "name": name,
            "category": data.get("category", "Thiết bị mạng LAN").strip(),
            "serial_number": (data.get("serial_number") or data.get("mac") or "").strip() or None,
            "quantity": int(data.get("quantity", 1)),
            "unit": data.get("unit", "Cái").strip() or "Cái",
            "location": data.get("location") or f"IP LAN: {data.get('ip', 'N/A')}",
            "site_name": data.get("site_name", "Trụ sở chính").strip() or "Trụ sở chính",
            "status": data.get("status", "Đang sử dụng").strip(),
            "warranty_end": data.get("warranty_end") or None,
            "notes": data.get("notes") or f"Phát hiện tự động từ Quét LAN (IP: {data.get('ip', '')})"
        }
        new_id = models.create_item(item_payload)
        return jsonify({"success": True, "id": new_id, "message": f"Đã thêm thiết bị '{name}' vào cơ sở '{item_payload['site_name']}'!"}), 201
    except Exception as e:
        return jsonify({"error": str(e)}), 400

# --- API ITEMS (VẬT TƯ / THIẾT BỊ) ---
@app.route("/api/items", methods=["GET"])
def api_get_items():
    search = request.args.get("q", "").strip()
    category = request.args.get("category", "").strip()
    status = request.args.get("status", "").strip()
    site_name = request.args.get("site", "").strip() or request.args.get("site_name", "").strip()
    items = models.get_all_items(search=search, category=category, status=status, site_name=site_name)
    return jsonify(items)

@app.route("/api/items/<int:item_id>", methods=["GET"])
def api_get_item(item_id):
    item = models.get_item_by_id(item_id)
    if not item:
        return jsonify({"error": "Không tìm thấy thiết bị"}), 404
    return jsonify(item)

@app.route("/api/items", methods=["POST"])
def api_create_item():
    data = request.json or {}
    if not data.get("name") or not data.get("category"):
        return jsonify({"error": "Tên thiết bị và Loại thiết bị là bắt buộc"}), 400
    try:
        new_id = models.create_item(data)
        return jsonify({"success": True, "id": new_id, "message": "Thêm thiết bị mới thành công!"}), 201
    except Exception as e:
        return jsonify({"error": str(e)}), 400

@app.route("/api/items/<int:item_id>", methods=["PUT"])
def api_update_item(item_id):
    data = request.json or {}
    if not data.get("name") or not data.get("category"):
        return jsonify({"error": "Tên thiết bị và Loại thiết bị là bắt buộc"}), 400
    try:
        models.update_item(item_id, data)
        return jsonify({"success": True, "message": "Cập nhật thông tin thành công!"})
    except Exception as e:
        return jsonify({"error": str(e)}), 400

@app.route("/api/items/<int:item_id>", methods=["DELETE"])
def api_delete_item(item_id):
    try:
        models.delete_item(item_id)
        return jsonify({"success": True, "message": "Đã xóa thiết bị thành công!"})
    except Exception as e:
        return jsonify({"error": str(e)}), 400

# --- API TRANSACTIONS (NHẬP / XUẤT KHO) ---
@app.route("/api/transactions", methods=["GET"])
def api_get_transactions():
    search = request.args.get("q", "").strip()
    tx_type = request.args.get("type", "").strip()
    txs = models.get_all_transactions(search=search, tx_type=tx_type)
    return jsonify(txs)

@app.route("/api/transactions", methods=["POST"])
def api_create_transaction():
    data = request.json or {}
    if not data.get("item_id") or not data.get("quantity"):
        return jsonify({"error": "Vui lòng chọn thiết bị và nhập số lượng"}), 400
    try:
        new_tx_id = models.create_transaction(data)
        return jsonify({"success": True, "id": new_tx_id, "message": "Giao dịch kho đã hoàn tất thành công!"}), 201
    except ValueError as ve:
        return jsonify({"error": str(ve)}), 400
    except Exception as e:
        return jsonify({"error": str(e)}), 500

# --- XUẤT BÁO CÁO EXCEL ---
@app.route("/api/export/inventory", methods=["GET"])
def export_inventory_excel():
    """Tạo và tải về file Excel danh mục vật tư thiết bị bao gồm cơ sở / chi nhánh."""
    site_name = request.args.get("site", "").strip()
    items = models.get_all_items(site_name=site_name)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh Mục Vật Tư IT"

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    headers = [
        "STT", "Tên thiết bị / Vật tư", "Phân loại", "Cơ sở / Chi nhánh", "Serial Number (S/N)",
        "Số lượng", "Đơn vị", "Vị trí kho", "Trạng thái", "Hạn bảo hành", "Ghi chú", "Cập nhật lúc"
    ]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for i, it in enumerate(items, 1):
        row_data = [
            i,
            it["name"],
            it["category"],
            it.get("site_name") or "Trụ sở chính",
            it["serial_number"] or "—",
            it["quantity"],
            it["unit"],
            it["location"] or "—",
            it["status"],
            it["warranty_end"] or "—",
            it["notes"] or "",
            it["updated_at"]
        ]
        ws.append(row_data)

        for col_idx in range(1, len(row_data) + 1):
            c = ws.cell(row=i + 1, column=col_idx)
            c.border = thin_border
            if col_idx in [1, 6]:
                c.alignment = Alignment(horizontal="center")
            elif col_idx in [9, 10]:
                c.alignment = Alignment(horizontal="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    now_str = datetime.now().strftime("%Y%m%d_%H%M")
    filename = f"BaoCao_VatTu_IT_{now_str}.xlsx"

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )

@app.route("/api/export/transactions", methods=["GET"])
def export_transactions_excel():
    """Tạo và tải về file Excel lịch sử Nhập/Xuất kho."""
    txs = models.get_all_transactions()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lịch Sử Nhập Xuất"

    header_fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    headers = [
        "Mã GD", "Tên thiết bị / Vật tư", "Loại giao dịch", "Số lượng", "Đơn vị",
        "Người thực hiện", "Người nhận bàn giao", "Phòng ban / Khách hàng", "Ghi chú mục đích", "Thời gian"
    ]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for i, tx in enumerate(txs, 1):
        row_data = [
            f"GD{tx['id']:05d}",
            tx["item_name"],
            tx["type"],
            tx["quantity"],
            tx["unit"],
            tx["performer"],
            tx["recipient"] or "—",
            tx["department"] or "—",
            tx["notes"] or "",
            tx["created_at"]
        ]
        ws.append(row_data)

        for col_idx in range(1, len(row_data) + 1):
            c = ws.cell(row=i + 1, column=col_idx)
            c.border = thin_border
            if col_idx in [1, 3, 4, 5]:
                c.alignment = Alignment(horizontal="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    now_str = datetime.now().strftime("%Y%m%d_%H%M")
    filename = f"LichSu_NhapXuat_Kho_{now_str}.xlsx"

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )
import multiprocessing

if __name__ == "__main__":
    # Bắt buộc phải có để PyInstaller trên Windows không bị lặp vô tận
    multiprocessing.freeze_support()

    print(f"\n=======================================================")
    print(f"   ỨNG DỤNG QUẢN LÝ VẬT TƯ & THIẾT BỊ MẠNG IT         ")
    print(f"=======================================================")
    print(f" [*] Đang khởi chạy máy chủ Web:")
    print(f"     - Truy cập cục bộ : http://127.0.0.1:8000")
    print(f"     - Truy cập LAN    : http://{LOCAL_IP}:8000")
    print(f"=======================================================\n")
    app.run(host="0.0.0.0", port=8000, debug=False)
